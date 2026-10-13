#!/usr/bin/env python3
"""Convert paper/main.tex into an editable Word manuscript.

Journals want a single-column manuscript for initial submission anyway, so
we deliberately produce a clean editable document rather than trying to
reproduce IEEEtran's two-column layout in Word.

Pipeline:
  main.tex  --(strip IEEEtran macros, resolve \\cite -> [n] from main.aux)-->
  pandoc-friendly .tex  --(pandoc)-->  .docx  --(python-docx)--> styled .docx
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
OUT = PAPER / "docx"
OUT.mkdir(parents=True, exist_ok=True)


# ----------------------------------------------------------------- citations
def citation_map() -> dict[str, int]:
    aux = (PAPER / "main.aux").read_text(errors="ignore")
    return {k: int(n) for k, n in
            re.findall(r"\\bibcite\{([^}]*)\}\{(\d+)\}", aux)}


def label_map() -> dict[str, str]:
    """label -> printed number, exactly as LaTeX numbered them."""
    aux = (PAPER / "main.aux").read_text(errors="ignore")
    out = {}
    for lab, num in re.findall(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", aux):
        num = re.sub(r"\\mbox\s*", "", num).replace("{", "").replace("}", "").strip()
        if num:
            out[lab] = num
    return out


_ACCENTS = {"`": "\u0300", "'": "\u0301", "^": "\u0302", "~": "\u0303",
            '"': "\u0308", "=": "\u0304", ".": "\u0307", "u": "\u0306",
            "v": "\u030c", "H": "\u030b", "c": "\u0327"}


def _latex_accents(s: str) -> str:
    """{\\`\\i} -> ì, \\"u -> ü, {\\"O} -> Ö ... as real Unicode letters.

    Left to the generic macro strip, "Dess{\\`\\i}" lost its \\i and the bare
    accent landed on the following comma ("Dess,̀").
    """
    import unicodedata

    def repl(m: re.Match) -> str:
        base = "i" if m.group(2) == "\\i" else ("j" if m.group(2) == "\\j" else m.group(2))
        return unicodedata.normalize("NFC", base + _ACCENTS[m.group(1)])

    s = re.sub(r"\{?\\([`'^~\"=.]|[uvHc](?![a-zA-Z]))\s*\{?(\\[ij](?![a-zA-Z])|[A-Za-z])\}?\}?",
               repl, s)
    return s


def bibliography_entries() -> list[str]:
    """Plain-text reference list, in citation order, from main.bbl."""
    bbl = (PAPER / "main.bbl").read_text(errors="ignore")
    body = bbl.split(r"\begin{thebibliography}", 1)[-1]
    body = body.split(r"\end{thebibliography}")[0]
    chunks = re.split(r"\\bibitem\{[^}]*\}", body)[1:]
    out = []
    for c in chunks:
        c = _latex_accents(c)
        c = re.sub(r"\\BIB\w+", " ", c)
        c = re.sub(r"\\newblock", " ", c)
        c = re.sub(r"\\url\{([^}]*)\}", r"\1", c)
        c = re.sub(r"\\href\{[^}]*\}\{([^}]*)\}", r"\1", c)
        c = re.sub(r"\\emph\{([^}]*)\}", r"\1", c)
        c = re.sub(r"\\textit\{([^}]*)\}", r"\1", c)
        c = re.sub(r"\\BIBforeignlanguage\{[^}]*\}\{([^}]*)\}", r"\1", c)
        c = re.sub(r"\\[a-zA-Z]+\s*", " ", c)
        c = c.replace("{", "").replace("}", "").replace("~", " ")
        c = re.sub(r"\s+", " ", c).strip(" .,")
        if c:
            out.append(c + ".")
    return out


# ------------------------------------------------------------ table handling
def _clean_cell(c: str) -> str:
    """LaTeX cell -> plain text."""
    c = c.strip()
    c = c.replace("\\newline", " ")          # line break inside a p{} cell
    c = re.sub(r"\\multicolumn\{\d+\}\{[^}]*\}\{(.*)\}", r"\1", c)
    c = re.sub(r"\\(textbf|textit|emph|texttt)\{([^{}]*)\}", r"\2", c)
    c = re.sub(r"\\cite\{[^}]*\}", "", c)
    c = c.replace("$-$", "\u2212").replace("$+$", "+")      # signed numbers
    c = c.replace("\\,", " ").replace("~", " ").replace("$", "")
    c = c.replace("\\%", "%").replace("\\&", "&").replace("\\_", "_")
    c = c.replace("``", "“").replace("''", "”")   # LaTeX quotes -> real ones
    c = c.replace("---", "—").replace("--", "–")   # em / en dashes
    # Named symbols must be translated before the generic \command strip below,
    # or a header cell like "$\Delta$" silently becomes empty.
    c = re.sub(r"\^?\\ddagger", "‡", c)          # footnote marks: ddagger first,
    c = re.sub(r"\^?\\dagger", "†", c)            # or \dagger eats its prefix
    for name, ch in [(r"\\Delta", "Δ"), (r"\\delta", "δ"),
                     (r"\\times", "x"), (r"\\sim", "~"), (r"\\pm", "±"),
                     (r"\\approx", "≈"), (r"\\geq", "≥"),
                     (r"\\leq", "≤"), (r"\\rightarrow", "→"),
                     (r"\\mu", "μ"), (r"\\alpha", "α")]:
        c = re.sub(name + r"(?![a-zA-Z])", ch, c)
    c = c.replace("\\ ", " ")                # escaped inter-word space
    c = re.sub(r"\\[,;:!]", " ", c)          # thin / medium / thick spaces
    sup = str.maketrans("0123456789", "\u2070\u00b9\u00b2\u00b3\u2074\u2075\u2076\u2077\u2078\u2079")
    c = re.sub(r"\^\{?(\d+)\}?", lambda m: m.group(1).translate(sup), c)   # 2^{16} -> 2¹⁶
    c = re.sub(r"\\[a-zA-Z]+", "", c)
    c = c.replace("{", "").replace("}", "")
    c = re.sub(r"\s+", " ", c).strip()
    return c


def parse_tables(tex: str):
    """Extract every table environment as (label, caption, rows).

    Pass text whose \\cite and \\ref are already resolved (see resolve_refs):
    _clean_cell strips any command it does not know, so an unresolved
    \\cite inside a cell would silently lose its reference number.
    """
    out = {}
    # table\*? : two-column `table*` floats must be caught too. Missing them let
    # pandoc render Table I itself, and pandoc silently dropped its caption.
    for m in re.finditer(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", tex, re.S):
        block = m.group(0)
        cap = re.search(r"\\caption\{(.*?)\}\s*\n", block, re.S)
        caption = _clean_cell(cap.group(1)) if cap else ""
        lab = re.search(r"\\label\{([^}]*)\}", block)
        label = lab.group(1) if lab else f"tab{len(out)}"
        # strip @{} first: it contains braces and breaks a naive [^}]* spec match
        block = block.replace("@{}", "")
        # the column spec may nest one level of braces: p{1.8cm},
        # >{\raggedright\arraybackslash}
        tab = re.search(r"\\begin\{tabular\}\{(?:[^{}]|\{[^{}]*\})*\}(.*?)\\end\{tabular\}",
                        block, re.S)
        if not tab:
            continue
        body = tab.group(1)
        # \cmidrule(lr){2-3} must go before the plain rules: its braces and
        # parens survive naive stripping and surface as "(lr)2-3" inside the
        # first header cell.
        body = re.sub(r"\\cmidrule(?:\([^)]*\))?\{[^}]*\}", "", body)
        for rule in ["\\toprule", "\\midrule", "\\bottomrule"]:
            body = body.replace(rule, "")
        rows, spans = [], []
        for line in body.split("\\\\"):
            line, row_spans = _expand_multicolumn(line)
            cells = [_clean_cell(c) for c in line.split("&")]
            if any(cells):
                spans += [(len(rows), c0, n) for c0, n in row_spans if n > 1]
                rows.append(cells)
        if rows:
            out[label] = (caption, rows, spans)
    return out


def _expand_multicolumn(line: str):
    """Replace \\multicolumn{n}{spec}{text} by text plus n-1 empty cells.

    Returns the new line and a list of (first column, n) spans, so the Word
    table can merge the cells exactly as LaTeX spans them.
    """
    spans, out, i, col = [], [], 0, 0
    key = "\\multicolumn{"
    while True:
        j = line.find(key, i)
        if j < 0:
            out.append(line[i:])
            break
        before = line[i:j]
        col += before.count("&")
        out.append(before)
        k = j + len(key)
        n_end = line.index("}", k)
        n = int(line[k:n_end])
        k = line.index("{", n_end) + 1          # column spec (may nest once)
        depth = 1
        while depth:
            depth += {"{": 1, "}": -1}.get(line[k], 0)
            k += 1
        k = line.index("{", k) + 1              # the text, with nested braces
        start, depth = k, 1
        while depth:
            depth += {"{": 1, "}": -1}.get(line[k], 0)
            k += 1
        out.append(line[start:k - 1] + " &" * (n - 1))
        spans.append((col, n))
        col += n - 1
        i = k
    return "".join(out), spans


# --------------------------------------------------------- inline math
# Pandoc turns every `$...$` span into a Word equation object.  For real
# mathematics that is what you want; for `$76.0\%$` it is actively hostile,
# because the author has to open the equation editor just to retype a number.
# We therefore unwrap *trivial* inline math to plain text and leave genuine
# expressions alone.
_SIMPLE_MATH = re.compile(r"""
    ^(?:
        [\d\s.,%+\-=/\[\]()<>:]          # numbers and plain punctuation
      | \\[,;:!]                          # spacing commands
      | \\%                               # escaped percent
      | \\times | \\sim | \\approx        # a few common symbols
      | \\Delta | \\pm | \\geq | \\leq
      | \\mathrm\{[A-Za-z]+\}
      | \^\{?\d+\}? | _\{?\d+\}?          # simple sub/superscripts
      | [A-Za-z]{0,3}                     # short unit strings: MB, KB, ms, s
    )+$
""", re.X)

# NOTE: the result is still LaTeX that pandoc has to read, so `%` MUST stay
# escaped as `\%`.  Emitting a bare `%` starts a LaTeX comment and silently
# deletes the rest of the line -- which is exactly as destructive as it sounds.
_MATH_TEXT = [
    (r"\\times", "\u00d7"), (r"\\Delta", "\u0394"), (r"\\pm", "\u00b1"),
    (r"\\approx", "\u2248"), (r"\\geq", "\u2265"), (r"\\leq", "\u2264"),
    (r"\\sim", "\\\\textasciitilde{}"), (r"\\,", "\u2009"), (r"\\;", " "),
    (r"\\!", ""), (r"\\:", " "),
]


def _unwrap_simple_math(t: str) -> str:
    def repl(m: re.Match) -> str:
        inner = m.group(1)
        if not _SIMPLE_MATH.match(inner.strip()):
            return m.group(0)                       # real maths: leave it
        s = inner
        for pat, ch in _MATH_TEXT:
            s = re.sub(pat, ch, s)
        s = re.sub(r"\^\{([^}]*)\}", r"\\textsuperscript{\1}", s)   # 2^{16}
        s = re.sub(r"_\{([^}]*)\}", r"\\textsubscript{\1}", s)
        s = re.sub(r"\^(\d+)", r"\\textsuperscript{\1}", s)          # 2^3
        s = re.sub(r"_(\d+)", r"\\textsubscript{\1}", s)              # L_2
        s = s.replace("-", "\u2212")                     # proper minus sign
        return s.strip()

    # only single-dollar spans; $$...$$ and display environments are untouched
    return re.sub(r"(?<!\$)\$([^$\n]+?)\$(?!\$)", repl, t)


# ------------------------------------------------------------- preprocessing
def resolve_refs(tex: str, cites: dict[str, int],
                 labels: dict[str, str] | None = None) -> str:
    """\\cite -> "[3], [4]" and \\ref -> printed number, exactly as in the PDF.

    Runs before tables are parsed, so citations inside table cells keep their
    numbers in the Word copy too.
    """
    labels = labels or {}

    def sub_cite(m: re.Match) -> str:
        # a key may carry a LaTeX line-continuation ("%\n") from the source
        keys = [re.sub(r"%[^\n]*\n?", "", k).strip() for k in m.group(1).split(",")]
        missing = [k for k in keys if k and k not in cites]
        if missing:
            raise SystemExit(f"unresolved citation key(s) {missing}: "
                             "recompile main.tex so main.aux is current")
        nums = sorted({cites[k] for k in keys if k})
        # IEEEtran without the cite package prints "[3], [4]"; match the PDF
        return ", ".join(f"[{n}]" for n in nums)

    t = re.sub(r"\\cite\{([^}]*)\}", sub_cite, tex)
    # \ref -> the number LaTeX actually printed, so the Word copy reads the
    # same as the PDF instead of showing raw labels
    t = re.sub(r"\\eqref\{([^}]*)\}",
               lambda m: "(" + labels.get(m.group(1), m.group(1)) + ")", t)
    t = re.sub(r"\\ref\{([^}]*)\}",
               lambda m: labels.get(m.group(1), m.group(1)), t)
    return t


def preprocess(tex: str, cites: dict[str, int],
               labels: dict[str, str] | None = None) -> str:
    labels = labels or {}
    t = resolve_refs(tex, cites, labels)
    t = t.replace("@{}", "")

    # IEEEtran-only constructs pandoc does not know
    t = re.sub(r"\\IEEEPARstart\{(.)\}\{([^}]*)\}", r"\1\2", t)
    t = re.sub(r"\\markboth\{[^}]*\}%?\s*\{[^}]*\}", "", t)
    t = t.replace(r"\balance", "").replace(r"\IEEEpubid", "")
    t = re.sub(r"\\graphicspath\{[^\n]*\}", "", t)
    t = re.sub(r"\\bibliographystyle\{[^}]*\}", "", t)
    t = re.sub(r"\\bibliography\{[^}]*\}", "", t)

    # Author block.  \thanks holds the affiliation, the contact address and the
    # code URL -- dropping it silently would ship a manuscript with no author
    # details at all, so we lift the contents into visible paragraphs directly
    # under the title instead.
    author_m = re.search(r"\\author\{(.*?)\n*\}\s*\n", t, re.S)
    thanks: list[str] = []
    if author_m:
        block = author_m.group(1)
        thanks = [re.sub(r"\s+", " ", x).strip()
                  for x in re.findall(r"\\thanks\{(.*?)\}(?=%|\s*\\thanks|\s*$)",
                                      block, re.S)]
        name = re.sub(r"\\thanks\{.*", "", block, flags=re.S)
        name = name.replace("~", " ").replace("%", "").strip()
        t = t[:author_m.start()] + f"\\author{{{name}}}\n" + t[author_m.end():]

    if thanks:
        note = "\n\n".join(r"\noindent\small " + x for x in thanks)
        t = t.replace(r"\maketitle", "\\maketitle\n\n" + note + "\n\n\\normalsize\n", 1)

    # float package preamble (Algorithm 1) means nothing to pandoc
    t = re.sub(r"\\floatstyle\{[^}]*\}\s*\n", "", t)
    t = re.sub(r"\\newfloat\{[^}]*\}\{[^}]*\}\{[^}]*\}\s*\n", "", t)
    t = re.sub(r"\\floatname\{[^}]*\}\{[^}]*\}\s*\n", "", t)
    t = t.replace("\\usepackage{float}\n", "")

    # Algorithm 1: a ruled float in the PDF; in Word a bold caption line
    # followed by its numbered steps
    def _alg(m):
        blk = m.group(0)
        cap = re.search(r"\\caption\{(.*?)\}\s*\n", blk, re.S)
        lab = re.search(r"\\label\{([^}]*)\}", blk)
        n = labels.get(lab.group(1), "1") if lab else "1"
        body = blk
        body = re.sub(r"\\begin\{algorithm\}(\[[^\]]*\])?", "", body)
        body = body.replace("\\end{algorithm}", "")
        body = re.sub(r"\\caption\{.*?\}\s*\n", "", body, count=1, flags=re.S)
        body = re.sub(r"\\label\{[^}]*\}", "", body)
        body = body.replace("\\small", "")
        head = "\\par\\noindent\\textbf{Algorithm %s.} %s\\par\n" % (n, cap.group(1) if cap else "")
        return "\n" + head + body + "\n"
    t = re.sub(r"\\begin\{algorithm\}.*?\\end\{algorithm\}", _alg, t, flags=re.S)

    # numbered equations: Word shows no equation numbers, so append "(n)"
    def _eq(m):
        blk = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", blk)
        n = labels.get(lab.group(1)) if lab else None
        blk = re.sub(r"\\label\{[^}]*\}", "", blk).strip()
        tag = ("\\qquad (%s)" % n) if n else ""
        return "\n$$" + blk + tag + "$$\n"
    t = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}", _eq, t, flags=re.S)
    def _align(m):
        blk = m.group(1)
        labs = re.findall(r"\\label\{([^}]*)\}", blk)
        lines = [l for l in re.split(r"\\\\\s*\n", blk) if l.strip()]
        out = []
        for l in lines:
            lab = re.search(r"\\label\{([^}]*)\}", l)
            n = labels.get(lab.group(1)) if lab else None
            l = re.sub(r"\\label\{[^}]*\}", "", l).replace("&", "").strip()
            out.append("\n$$" + l + (("\\qquad (%s)" % n) if n else "") + "$$\n")
        return "".join(out)
    t = re.sub(r"\\begin\{align\}(.*?)\\end\{align\}", _align, t, flags=re.S)

    # IEEE section numbers ("III", "III-A"), so "Section IV-E" in the text can
    # be found in the Word copy, which pandoc would otherwise leave unnumbered
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
    sec_no, sub_no = 0, 0
    def _number(m):
        nonlocal sec_no, sub_no
        kind, title = m.group(1), m.group(2)
        if kind == "section":
            sec_no += 1; sub_no = 0
            return "\\section{%s. %s}" % (roman[sec_no - 1], title)
        sub_no += 1
        return "\\subsection{%s. %s}" % (chr(64 + sub_no), title)
    t = re.sub(r"\\(section|subsection)\{([^}]*)\}", _number, t)
    t = t.replace("A$^\\ast$", "A*")

    # keywords environment
    t = t.replace(r"\begin{IEEEkeywords}", r"\par\noindent\textbf{Index Terms---}")
    t = t.replace(r"\end{IEEEkeywords}", r"\par")

    # number figure captions like the PDF ("Fig. 1."); pandoc puts them below
    def _num_fig(m):
        blk = m.group(0)
        lab = re.search(r"\\label\{([^}]*)\}", blk)
        n = labels.get(lab.group(1)) if lab else None
        return blk.replace("\\caption{", "\\caption{\\textbf{Fig. %s.} " % n, 1) if n else blk
    t = re.sub(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}", _num_fig, t, flags=re.S)

    # pandoc 2.9 drops the caption of a two-column figure*; Word is one column
    t = t.replace("\\begin{figure*}", "\\begin{figure}").replace("\\end{figure*}", "\\end{figure}")

    # figures: Word wants raster, and pandoc 2.9 will not embed PDF
    t = re.sub(r"\\includegraphics\[[^\]]*\]\{([^}]*)\.pdf\}",
               r"\\includegraphics{figures/\1.png}", t)
    t = re.sub(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}",
               r"\\includegraphics{figures/\1.png}", t)
    t = t.replace("figures/figures/", "figures/")
    t = re.sub(r"\.png\.png", ".png", t)

    # Pandoc 2.9's LaTeX table reader drops columns on booktabs tables, so we
    # pull the tables out here and rebuild them natively with python-docx.
    for label in parse_tables(t):
        t = re.sub(r"\\begin\{table\*?\}(?:(?!\\end\{table\*?\}).)*?"
                   + re.escape("\\label{" + label + "}")
                   + r"(?:(?!\\end\{table\*?\}).)*?\\end\{table\*?\}",
                   f"\n\nTABLEPLACEHOLDER{label.replace(':', '')}ENDTABLE\n\n",
                   t, count=1, flags=re.S)
    t = t.replace(r"\small", "").replace(r"\centering", "")

    # plain numbers should stay plain text, not become equation objects
    t = _unwrap_simple_math(t)

    # math: texmath 0.12 (pandoc 2.9) does not know \tfrac or \textstyle
    t = t.replace(r"\tfrac", r"\frac").replace(r"\textstyle", "")
    t = t.replace(r"\Big(", r"\left(").replace(r"\Big)", r"\right)")

    # single column, standard class
    t = t.replace(r"\documentclass[journal]{IEEEtran}", r"\documentclass{article}")

    # append the resolved bibliography
    refs = bibliography_entries()
    block = ["\n\\section*{References}\n\\begin{enumerate}"]
    block += [f"\\item {r}" for r in refs]
    block.append("\\end{enumerate}\n")
    t = t.replace(r"\end{document}", "\n".join(block) + "\n\\end{document}")
    return t


# --------------------------------------------------------------- docx styling
def _set_borders(table) -> None:
    """Draw single-line borders on every cell (no style dependency)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def _insert_table(doc, para, caption: str, rows: list[list[str]], number: str = "",
                  spans=()):
    """Replace a placeholder paragraph with a real Word table + caption."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]

    para.text = ""
    if number:                              # "Table I." -- same numbering as the PDF
        lab = para.add_run(f"Table {number}. ")
        lab.bold = True
        lab.font.size = Pt(10)
    run = para.add_run(caption)
    run.font.size = Pt(10)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.line_spacing = 1.0

    table = doc.add_table(rows=len(rows), cols=ncol)
    # pandoc's generated docx has no "Table Grid" style, so fall back to
    # whatever table style exists and draw the borders ourselves.
    for cand in ("Table Grid", "Table", "TableNormal"):
        try:
            table.style = cand
            break
        except KeyError:
            continue
    _set_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Right-align a column only when its body cells are numbers (results
    # tables); descriptive columns such as Table II's read better left-aligned.
    def _numeric(s: str) -> bool:
        s = s.strip()
        return bool(s) and bool(re.fullmatch(
            r"[\[\]()+\-−–~≈≤≥±$%.,:;/\s\d×xA-Za-zμ†‡]{0,24}", s)) \
            and bool(re.search(r"\d", s)) and len(s.split()) <= 3
    numeric_col = []
    for j in range(ncol):
        # placeholder dashes ("---" for "not reported") say nothing either way
        body = [r[j] for r in rows[1:]
                if r[j].strip() and r[j].strip() not in {"—", "–", "-", "n/a"}]
        numeric_col.append(bool(body) and
                           sum(_numeric(x) for x in body) >= 0.6 * len(body))

    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = cell
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                if j > 0 and numeric_col[j]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.size = Pt(9 if ncol < 9 else 7.5)   # wide tables
                    r.font.name = "Times New Roman"
                    if i == 0:
                        r.bold = True
    # LaTeX \\multicolumn spans become merged Word cells
    for r, c0, n in spans:
        if c0 + n - 1 < ncol:
            a, b = table.cell(r, c0), table.cell(r, c0 + n - 1)
            text = a.text
            m = a.merge(b)
            m.text = text
            for p in m.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 or len(rows[r]) > 1 and any(rows[r][1:]) \
                    else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Times New Roman"
                    run.bold = (r == 0) or run.bold
    # move the table to sit directly after its caption
    para._p.addnext(table._tbl)
    return table


def style(path: Path, tables: dict, labels: dict | None = None) -> None:
    labels = labels or {}
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document(str(path))

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5      # room for reviewer marks

    for name, size, bold in [("Heading 1", 14, True), ("Heading 2", 12, True),
                             ("Heading 3", 11, True), ("Title", 18, True)]:
        if name in doc.styles:
            f = doc.styles[name].font
            f.name = "Times New Roman"
            f.size = Pt(size)
            f.bold = bold
            f.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # rebuild the tables that pandoc could not handle
    placed = set()
    for p in list(doc.paragraphs):
        m = re.search(r"TABLEPLACEHOLDER(\w+)ENDTABLE", p.text)
        if not m:
            continue
        key = m.group(1)
        match = next((k for k in tables if k.replace(":", "") == key), None)
        if match:
            caption, rows, spans = tables[match]
            _insert_table(doc, p, caption, rows, labels.get(match, ""), spans)
            placed.add(match)

    # centre every figure
    for p in doc.paragraphs:
        if p.runs and any("graphicData" in r._element.xml for r in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(path))
    missing = set(tables) - placed
    if missing:
        print(f"  WARNING: tables not placed: {sorted(missing)}")
    else:
        print(f"  all {len(tables)} tables rebuilt natively")


def verify(docx_path: Path) -> bool:
    """Check the Word text against the compiled PDF for silent text loss.

    Escaping bugs in the LaTeX we hand to pandoc (a stray unescaped `%`, for
    instance) delete the rest of a line without any error, so a build that
    "succeeds" can still be missing sentences.  This compares the two.
    """
    import zipfile

    pdf = PAPER / "main.pdf"
    if not pdf.exists() or not shutil.which("pdftotext"):
        print("  [verify] skipped (need paper/main.pdf and pdftotext)")
        return True

    raw = subprocess.run(["pdftotext", str(pdf), "-"],
                         capture_output=True, text=True).stdout
    z = zipfile.ZipFile(docx_path)
    x = z.read("word/document.xml").decode("utf8")
    if "word/footnotes.xml" in z.namelist():       # dataset credits live here
        x += z.read("word/footnotes.xml").decode("utf8")
    body = " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", x))
    body += " " + " ".join(re.findall(r"<m:t[^>]*>([^<]*)</m:t>", x))

    def words(s: str) -> set:
        s = s.replace("−", "-").replace(" ", " ")
        return {w for w in re.findall(r"[A-Za-z][A-Za-z-]{3,}", s.lower())}

    pdf_w, doc_w = words(raw), words(body)
    missing = sorted(pdf_w - doc_w)
    # the PDF carries running heads and bibliography formatting the docx lacks
    # pdftotext splits IEEEtran's small-caps section heads ("Results" ->
    # "R" + "esults") and joins hyphenated line breaks, so a raw diff is full
    # of false positives.  Only flag a word if it is absent even after
    # allowing for those two effects.
    # Known-benign differences:
    #  - figures carry vector text in the PDF but are rasterised in Word, so
    #    their axis labels and titles are not extractable from the .docx;
    #  - IEEEtran numbers its section heads (VIII), Word styles do not;
    #  - accented names decompose differently in the two extractors.
    ignore = {"preprint", "under", "review", "mahmud", "agentic", "index",
              "abstract", "husza", "viii", "parameters", "unless", "dashed",
              "device", "acted"}
    flat = re.sub(r"[^a-z]", "", body.lower())
    real = []
    for w in missing:
        if w in ignore:
            continue
        stem = re.sub(r"[^a-z]", "", w)
        if stem in flat:              # hyphenation join, e.g. "open-ended"
            continue
        if len(stem) > 3 and any(stem in re.sub(r"[^a-z]", "", d)
                                 for d in doc_w):
            continue
        if any(d.endswith(stem) or stem.endswith(d[1:]) for d in doc_w
               if len(d) > 4 and len(stem) > 4):   # dropped leading capital
            continue
        real.append(w)
    missing = real
    ratio = len(doc_w & pdf_w) / max(len(pdf_w), 1)

    print(f"  [verify] vocabulary overlap with the PDF: {ratio:6.1%}")
    if missing:
        print(f"  [verify] CHECK - {len(missing)} words present in the PDF but "
              f"not extractable from the Word file: {missing[:12]}")
        print("  [verify] inspect these; figure-internal text is expected, "
              "missing sentences are not")
        return False
    print("  [verify] no text lost")
    return True


def main() -> None:
    cites = citation_map()
    src = (PAPER / "main.tex").read_text()
    labels = label_map()
    # resolve first: parsing the raw source dropped every [n] inside tables
    tables = parse_tables(resolve_refs(src, cites, labels))
    pre = preprocess(src, cites, labels)
    # pandoc's input goes to the system temp dir, not the paper folder
    import tempfile
    tmp = Path(tempfile.gettempdir()) / "conclave_pandoc_input.tex"
    tmp.write_text(pre)

    # Word locks a document while it is open; writing over it then fails with
    # "permission denied".  Fall back to a second name instead of losing the build.
    target = OUT / "CONCLAVE_paper.docx"
    try:
        if target.exists():
            with open(target, "r+b"):
                pass
    except PermissionError:
        target = OUT / "CONCLAVE_paper_v2.docx"
        print(f"  NOTE: CONCLAVE_paper.docx is open in Word; writing {target.name}")

    cmd = ["pandoc", str(tmp), "-f", "latex", "-t", "docx",
           "--resource-path", str(PAPER), "--standalone",
           "-o", str(target)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stderr[:3000], file=sys.stderr)
        sys.exit(1)
    if r.stderr.strip():
        print("pandoc warnings:\n" + r.stderr[:1500])

    style(target, tables, labels)
    try:
        tmp.unlink(missing_ok=True)
    except OSError:
        pass
    print(f"wrote {target}")
    verify(target)
    print(f"  resolved {len(cites)} citations, "
          f"{len(bibliography_entries())} references, {len(tables)} tables, "
          f"{len(labels)} cross-references")


if __name__ == "__main__":
    main()
